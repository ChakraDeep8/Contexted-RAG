from __future__ import annotations

import importlib.metadata
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import DocumentElement, ExtractedDocument, SourceLocation


class ExtractionError(RuntimeError):
    """Raised when a supported document cannot be extracted."""


class Extractor(Protocol):
    def extract(self, path: Path, relative_path: str, content_hash: str) -> ExtractedDocument: ...


class DoclingExtractor:
    """Docling-backed extractor with an XLSX provenance refinement pass."""

    def __init__(self) -> None:
        self._converter: Any = None

    @property
    def parser_version(self) -> str:
        try:
            return f"docling-{importlib.metadata.version('docling')}"
        except importlib.metadata.PackageNotFoundError:
            return "docling-unavailable"

    def _get_converter(self) -> Any:
        if self._converter is not None:
            return self._converter
        try:
            from docling.datamodel.base_models import InputFormat
            from docling.datamodel.pipeline_options import PdfPipelineOptions, RapidOcrOptions
            from docling.document_converter import DocumentConverter, ImageFormatOption, PdfFormatOption

            options = PdfPipelineOptions()
            options.do_ocr = True
            options.do_table_structure = True
            options.ocr_options = RapidOcrOptions()
            self._converter = DocumentConverter(
                format_options={
                    InputFormat.PDF: PdfFormatOption(pipeline_options=options),
                    InputFormat.IMAGE: ImageFormatOption(pipeline_options=options),
                }
            )
            return self._converter
        except Exception as exc:
            raise ExtractionError(
                "Docling/RapidOCR could not initialize. Install dependencies with "
                "'python -m pip install -e .'."
            ) from exc

    def extract(self, path: Path, relative_path: str, content_hash: str) -> ExtractedDocument:
        if path.suffix.lower() == ".xlsx":
            return self._extract_xlsx(path, relative_path, content_hash)
        return self._extract_docling(path, relative_path, content_hash)

    def _convert(self, path: Path) -> Any:
        try:
            result = self._get_converter().convert(path)
            status = str(getattr(result, "status", "success")).lower()
            if "failure" in status:
                raise ExtractionError(f"Docling conversion failed with status {status}")
            return result
        except ExtractionError:
            raise
        except Exception as exc:
            message = str(exc).strip() or exc.__class__.__name__
            raise ExtractionError(f"Could not extract {path.name}: {message}") from exc

    def _extract_docling(self, path: Path, relative_path: str, content_hash: str) -> ExtractedDocument:
        result = self._convert(path)
        document = result.document
        elements: list[DocumentElement] = []
        headings: list[str] = []

        try:
            iterator = document.iterate_items()
        except Exception as exc:
            raise ExtractionError(f"Docling document traversal failed: {exc}") from exc

        for item, level in iterator:
            label = str(getattr(item, "label", "text")).split(".")[-1].lower()
            text = self._item_text(item, document)
            if not text.strip():
                continue
            location = self._source_location(item)
            if "section_header" in label or label in {"title", "heading"}:
                depth = max(1, int(level or 1))
                headings = headings[: depth - 1] + [text.strip()]
                location.section = " > ".join(headings)
                continue
            if location.section is None and headings:
                location.section = " > ".join(headings)
            kind = "table" if "table" in label else ("list_item" if "list" in label else "text")
            elements.append(
                DocumentElement(kind=kind, text=text, location=location, heading_path=tuple(headings))
            )

        warnings = [str(error) for error in getattr(result, "errors", [])]
        raw = self._raw_document(document)
        if not elements:
            markdown = str(document.export_to_markdown()).strip()
            if markdown:
                elements.append(
                    DocumentElement(
                        kind="text", text=markdown,
                        location=SourceLocation(page=1 if path.suffix.lower() != ".pdf" else None),
                    )
                )
                warnings.append("Fine-grained provenance was unavailable; stored document-level text.")
        if not elements:
            raise ExtractionError("No extractable text or tables were found")
        return ExtractedDocument(
            path=path, relative_path=relative_path, content_hash=content_hash,
            title=path.stem, format=path.suffix.lower().lstrip("."), elements=elements,
            warnings=warnings, parser_version=self.parser_version, raw_document=raw,
        )

    def _extract_xlsx(self, path: Path, relative_path: str, content_hash: str) -> ExtractedDocument:
        result = self._convert(path)
        raw = self._raw_document(result.document)
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
        except Exception as exc:
            raise ExtractionError(f"Could not read workbook: {exc}") from exc

        elements: list[DocumentElement] = []
        warnings: list[str] = [str(error) for error in getattr(result, "errors", [])]
        try:
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                nonempty = [index for index, row in enumerate(rows, 1) if any(value is not None for value in row)]
                if not nonempty:
                    warnings.append(f"Sheet '{sheet.title}' is blank and was skipped.")
                    continue
                first_row = nonempty[0]
                max_column = max(
                    (index for row in rows for index, value in enumerate(row, 1) if value is not None),
                    default=1,
                )
                header_values = rows[first_row - 1][:max_column]
                header = [self._cell_text(value) or f"Column {i}" for i, value in enumerate(header_values, 1)]
                for row_number in nonempty:
                    values = rows[row_number - 1][:max_column]
                    if row_number == first_row:
                        text = " | ".join(header)
                        role = "header"
                    else:
                        cells = [f"{header[i]}: {self._cell_text(value)}" for i, value in enumerate(values) if value is not None]
                        text = " | ".join(cells)
                        role = "data"
                    if not text:
                        continue
                    elements.append(
                        DocumentElement(
                            kind="table_row", text=text, heading_path=(sheet.title,),
                            location=SourceLocation(
                                sheet=sheet.title, section=sheet.title,
                                row_start=row_number, row_end=row_number,
                                cell_range=f"A{row_number}:{get_column_letter(max_column)}{row_number}",
                            ),
                            metadata={"role": role, "header": header},
                        )
                    )
        finally:
            workbook.close()
        if not elements:
            raise ExtractionError("Workbook contains no non-empty cells")
        return ExtractedDocument(
            path=path, relative_path=relative_path, content_hash=content_hash, title=path.stem,
            format="xlsx", elements=elements, warnings=warnings,
            parser_version=f"{self.parser_version}+openpyxl-{importlib.metadata.version('openpyxl')}",
            raw_document=raw,
        )

    @staticmethod
    def _cell_text(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _raw_document(document: Any) -> dict[str, Any]:
        try:
            value = document.export_to_dict()
            return value if isinstance(value, dict) else {"document": value}
        except Exception:
            return {"markdown": str(document.export_to_markdown())}

    @staticmethod
    def _item_text(item: Any, document: Any) -> str:
        text = getattr(item, "text", None)
        if text:
            return str(text).strip()
        for method_name in ("export_to_markdown", "export_to_html"):
            method = getattr(item, method_name, None)
            if not method:
                continue
            try:
                return str(method(doc=document)).strip()
            except TypeError:
                try:
                    return str(method()).strip()
                except Exception:
                    pass
            except Exception:
                pass
        return ""

    @staticmethod
    def _source_location(item: Any) -> SourceLocation:
        provenance = getattr(item, "prov", None) or []
        if not provenance:
            return SourceLocation()
        prov = provenance[0]
        bbox_obj = getattr(prov, "bbox", None)
        bbox = None
        if bbox_obj is not None:
            bbox = {
                key: float(getattr(bbox_obj, key))
                for key in ("l", "t", "r", "b")
                if getattr(bbox_obj, key, None) is not None
            } or None
        return SourceLocation(page=getattr(prov, "page_no", None), bbox=bbox)

